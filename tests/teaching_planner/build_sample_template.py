from pathlib import Path
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).parent / "fixtures" / "sample_scheme_template.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)

doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width, section.page_height = Inches(11), Inches(8.5)
section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.5)
section.right_margin = Inches(0.5)
section.header_distance = Inches(0.3)
section.footer_distance = Inches(0.3)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(9)
normal.paragraph_format.space_after = Pt(3)

header = section.header.paragraphs[0]
header.text = "WUC | Academic Quality Assurance | Controlled Teaching Document"
header.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in header.runs:
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(90, 82, 110)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(4)
r = title.add_run("SCHEME OF WORK")
r.bold = True
r.font.name = "Calibri"
r.font.size = Pt(18)
r.font.color.rgb = RGBColor(111, 66, 193)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(10)
r = subtitle.add_run("{{institution_name}} | {{department_name}}")
r.bold = True
r.font.size = Pt(10)

metadata = doc.add_table(rows=3, cols=4)
metadata.alignment = WD_TABLE_ALIGNMENT.CENTER
metadata.autofit = False
metadata.style = "Table Grid"
meta_widths = [Inches(1.1), Inches(3.35), Inches(1.15), Inches(4.4)]
meta_values = [
    ("Programme", "{{programme_name}}", "Academic year", "{{academic_year}}"),
    ("Course", "{{course_code}} — {{course_name}}", "Academic period", "{{academic_period}}"),
    ("Lecturer", "{{lecturer_name}}", "Document", "{{document_number}} | v{{document_version}} | {{approval_status}}"),
]
for row, values in zip(metadata.rows, meta_values):
    for idx, (cell, value) in enumerate(zip(row.cells, values)):
        cell.width = meta_widths[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.text = value
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        for run in cell.paragraphs[0].runs:
            run.font.size = Pt(8.5)
            if idx in (0, 2):
                run.bold = True
                run.font.color.rgb = RGBColor(76, 44, 139)
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_mar = tc_pr.first_child_found_in("w:tcMar")
        if tc_mar is None:
            tc_mar = OxmlElement("w:tcMar")
            tc_pr.append(tc_mar)
        for side, value_dxa in (("top", 80), ("bottom", 80), ("start", 100), ("end", 100)):
            node = OxmlElement(f"w:{side}")
            node.set(qn("w:w"), str(value_dxa))
            node.set(qn("w:type"), "dxa")
            tc_mar.append(node)

doc.add_paragraph().paragraph_format.space_after = Pt(2)
headers = ["Week / date", "Topic & subtopics", "Learning outcome", "Methods & activities", "Resources", "Assessment", "Hours", "Remarks"]
placeholders = [
    "W{{week_number}}\n{{session_date}}\n{{start_time}}–{{end_time}}",
    "{{topic}}\n{{subtopics}}",
    "{{learning_outcomes}}",
    "{{teaching_methods}}\nLecturer: {{lecturer_activities}}\nLearners: {{learner_activities}}",
    "{{resources}}",
    "{{assessment_method}}",
    "{{duration}}",
    "{{remarks}}",
]
widths = [Inches(0.8), Inches(2.0), Inches(1.55), Inches(2.15), Inches(1.15), Inches(1.15), Inches(0.55), Inches(0.65)]
table = doc.add_table(rows=2, cols=len(headers))
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
table.style = "Table Grid"
for row_index, values in enumerate((headers, placeholders)):
    row = table.rows[row_index]
    for idx, (cell, text) in enumerate(zip(row.cells, values)):
        cell.width = widths[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.text = text
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if row_index == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(7.5 if row_index else 8)
                run.bold = row_index == 0
                if row_index == 0:
                    run.font.color.rgb = RGBColor(255, 255, 255)
        if row_index == 0:
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "6F42C1")
            cell._tc.get_or_add_tcPr().append(shading)
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_w = tc_pr.first_child_found_in("w:tcW")
        tc_w.set(qn("w:w"), str(int(widths[idx].inches * 1440)))
        tc_w.set(qn("w:type"), "dxa")

tbl_pr = table._tbl.tblPr
tbl_w = tbl_pr.first_child_found_in("w:tblW")
tbl_w.set(qn("w:w"), "14400")
tbl_w.set(qn("w:type"), "dxa")
tbl_layout = OxmlElement("w:tblLayout")
tbl_layout.set(qn("w:type"), "fixed")
tbl_pr.append(tbl_layout)
header_row_pr = table.rows[0]._tr.get_or_add_trPr()
repeat = OxmlElement("w:tblHeader")
repeat.set(qn("w:val"), "true")
header_row_pr.append(repeat)

note = doc.add_paragraph()
note.paragraph_format.space_before = Pt(6)
note.paragraph_format.space_after = Pt(0)
r = note.add_run("References: {{references}}  |  Template preview watermark: {{watermark}}")
r.font.size = Pt(8)
r.font.color.rgb = RGBColor(90, 82, 110)

footer = section.footer.paragraphs[0]
footer.text = "Template-controlled layout • Approved syllabus content • Deterministic scheduling"
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in footer.runs:
    run.font.size = Pt(7.5)
    run.font.color.rgb = RGBColor(110, 110, 110)

doc.save(OUT)
print(OUT)

